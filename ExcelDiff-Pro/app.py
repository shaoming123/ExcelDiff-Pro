"""ExcelDiff Pro | 双语 Excel / CSV 对账与差异检查工具.

Run: streamlit run app.py
Everything in the source spreadsheets is processed in the session's server memory.
No database, cloud AI, spreadsheet telemetry, or persistence integrations.
"""
from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
import hashlib
import html
from io import StringIO
from pathlib import Path
from urllib.parse import quote
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from exceldiff import (
    CompareConfig, DataImportError, DataTable, column_options,
    compare_tables, inspect_table, list_sheets, load_table, mapping_diagnostics,
)
from exceldiff.core import (
    MATCHED, DIFFERENT, ONLY_A, ONLY_B, DUPLICATE,
    INVALID_KEY_A, INVALID_KEY_B, INVALID_VALUE,
)
from exceldiff.io import present
from exceldiff.presets import export_profile, load_profile
from exceldiff.report import export_report, safe_cell
from exceldiff.sample import demo_bytes

st.set_page_config(
    page_title="ExcelDiff Pro | Excel comparison & reconciliation",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)
# The CSS asset must contain CSS declarations ONLY (no <style> tags).
# st.html(Path) injects a stylesheet without rendering CSS as Markdown text.
# This fixes visible stray CSS that appeared in v1.0 after a premature </style>.
STYLE_FILE = Path(__file__).parent / "assets" / "theme.css"
st.html(STYLE_FILE)

CONTACT_EMAIL = "greatfoodprefer@gmail.com"


def init_state() -> None:
    defaults = {
        "demo_mode": False, "upload_epoch": 0, "preset_epoch": 0,
        "preset_config": None, "preset_sheets": None,
        "tables": {}, "report_signature": None,
        "report_result": None, "report_blob": None,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def reset_state() -> None:
    for key in ("demo_mode", "preset_config", "preset_sheets", "report_signature", "report_result", "report_blob"):
        st.session_state[key] = False if key == "demo_mode" else None
    st.session_state.upload_epoch += 1
    st.session_state.preset_epoch += 1
    st.session_state.tables = {}


def short_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def cache_table(content: bytes, filename: str, sheet: str) -> DataTable:
    token = hashlib.sha256(content).hexdigest() + ":" + filename + ":" + sheet
    cache: dict[str, DataTable] = st.session_state.tables
    if token not in cache:
        if len(cache) >= 4:
            cache.clear()
        cache[token] = load_table(content, filename, sheet)
    return cache[token]


def raw_preview(table: DataTable, limit: int = 7) -> pd.DataFrame:
    width = min(table.width, 12)
    items = []
    for row_number, raw in enumerate(table.rows[:limit], start=1):
        item = {"ROW": row_number}
        for index in range(width):
            item[get_column_letter(index + 1)] = present(raw[index], 45)
        items.append(item)
    return pd.DataFrame(items)


def csv_export(items: list[dict[str, Any]]) -> bytes:
    if not items:
        return "Key,Result\n".encode("utf-8-sig")
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=list(items[0]), lineterminator="\r\n")
    writer.writeheader()
    for record in items:
        writer.writerow({key: safe_cell(value) for key, value in record.items()})
    return output.getvalue().encode("utf-8-sig")


init_state()
locale = st.sidebar.selectbox("Language / 语言", options=["English", "简体中文"], index=0, key="ui_lang")
zh = locale == "简体中文"


def t(english: str, chinese: str) -> str:
    return chinese if zh else english


def render_contact_cta(*, compact: bool = False) -> None:
    """Display a bilingual, personal customization CTA (no tracking or forms)."""
    subject = t("Custom automation enquiry | ExcelDiff Pro", "ExcelDiff Pro 定制自动化咨询")
    mailto = f"mailto:{CONTACT_EMAIL}?subject={quote(subject)}"
    link = html.escape(mailto, quote=True)
    if compact:
        st.markdown(
            '<div class="contact-mini">'
            '<div class="contact-mini-kicker">' + html.escape(t("WORK WITH ME", "定制开发服务")) + '</div>'
            '<strong>' + html.escape(t("Have a more complex workflow?", "需要更复杂的自动化？")) + '</strong>'
            '<p>' + html.escape(t(
                "Custom Excel reconciliation, Python automation and business integrations.",
                "提供 Excel 专属对账、Python 自动化和业务系统集成开发。",
            )) + '</p>'
            f'<a href="{link}">' + html.escape(t("Email me ↗", "邮件联系 ↗")) + '</a>'
            f'<div class="contact-email">{html.escape(CONTACT_EMAIL)}</div>'
            '</div>',
            unsafe_allow_html=True,
        )
        return
    st.markdown(
        '<section class="contact-panel">'
        '<div class="contact-panel-kicker">' + html.escape(t("CUSTOM SOLUTIONS · FREELANCE DEVELOPMENT", "专属解决方案 · 自由开发服务")) + '</div>'
        '<h2>' + html.escape(t(
            "Need something beyond a simple spreadsheet comparison?",
            "需要比基本 Excel 对比更进一步的解决方案？",
        )) + '</h2>'
        '<p>' + html.escape(t(
            "Every business has its own file formats, matching rules and reporting requirements. "
            "If you need multi-file reconciliation, custom validation logic, scheduled Excel reports, "
            "Python / SQL automation, API integrations or a private internal tool, I can help you "
            "design and build a solution around your workflow.",
            "每家公司都有不同的 Excel 格式、对账规则与报表需求。如果你需要多文件对账、专属验证规则、"
            "定时 Excel 报告、Python／SQL 自动化、API 整合或内部专用工具，欢迎联系我讨论定制开发。",
        )) + '</p>'
        '<div class="contact-panel-actions">'
        f'<a class="contact-button" href="{link}">' + html.escape(t("Let’s discuss your project ↗", "邮件咨询定制项目 ↗")) + '</a>'
        f'<span class="contact-address">{html.escape(CONTACT_EMAIL)}</span>'
        '</div>'
        '<div class="contact-panel-hint">' + html.escape(t(
            "To get started, send a short description of your workflow and a sample layout using dummy data — no confidential files needed.",
            "联系时可以简单说明现有工作流程及预期输出，并提供虚构示例格式，无需发送机密文件。",
        )) + '</div>'
        '</section>',
        unsafe_allow_html=True,
    )


def section(kicker_en: str, kicker_zh: str, title_en: str, title_zh: str, detail_en: str = "", detail_zh: str = "") -> None:
    st.markdown(
        '<div class="section-intro">' + html.escape(t(kicker_en, kicker_zh)) + '</div>'
        '<div class="section-title">' + html.escape(t(title_en, title_zh)) + '</div>'
        '<div class="section-desc">' + html.escape(t(detail_en, detail_zh)) + '</div>',
        unsafe_allow_html=True,
    )


def card(title_en: str, title_zh: str, body_en: str, body_zh: str) -> None:
    st.markdown(
        '<div class="infocard"><strong>' + html.escape(t(title_en, title_zh)) + '</strong><p>'
        + html.escape(t(body_en, body_zh)) + '</p></div>', unsafe_allow_html=True,
    )


def choose_column(table: DataTable, label: str, initial: int, widget_key: str,
                  data_start: int, exclude: tuple[int, ...] = ()) -> int:
    labels = column_options(table, data_start)
    choices = [c for c in labels if c not in exclude]
    if not choices:
        st.error(t("No columns remain to select.", "没有可选的剩余栏位。"))
        st.stop()
    index = choices.index(initial) if initial in choices else 0
    return st.selectbox(label, options=choices, index=index,
                        format_func=lambda col: labels[col], key=widget_key)


def preview_keys(table: DataTable, cols: tuple[int, ...], start_row: int, count: int = 5) -> pd.DataFrame:
    items = []
    for row_number, raw in table.data_rows(start_row)[:count]:
        items.append({"ROW": row_number, **{get_column_letter(c + 1): present(raw[c], 55) for c in cols}})
    return pd.DataFrame(items)


def result_table(records: list[dict[str, Any]]) -> pd.DataFrame:
    if not records:
        return pd.DataFrame()
    frame = pd.DataFrame(records)
    if zh:
        status_names = {
            MATCHED: "一致", DIFFERENT: "不一致", ONLY_A: "仅在 A", ONLY_B: "仅在 B",
            DUPLICATE: "重复 Key", INVALID_KEY_A: "A 的 Key 为空", INVALID_KEY_B: "B 的 Key 为空",
            INVALID_VALUE: "金额无效",
        }
        frame["Result"] = frame["Result"].map(status_names).fillna(frame["Result"])
        frame = frame.rename(columns={
            "Key": "匹配键", "Result": "结果", "A · Excel rows": "A 原始行号",
            "B · Excel rows": "B 原始行号", "A · Value": "A 数值",
            "B · Value": "B 数值", "Difference (A - B)": "差额 (A-B)",
        })
    return frame


with st.sidebar:
    st.markdown('<div class="brand"><div class="brand-icon">≋</div><div><h3>ExcelDiff Pro</h3><small>RECONCILIATION STUDIO</small></div></div>', unsafe_allow_html=True)
    st.caption(t("Compare workbooks confidently. No database required.", "轻松比较 Excel，不需要数据库。"))
    st.divider()
    st.markdown("**" + t("WORKSPACE", "工作空间") + "**")
    st.markdown(t("🟢  In-memory processing", "🟢  服务器内存处理"))
    st.markdown(t("🔒  No account / database", "🔒  无账号 / 无数据库"))
    st.markdown(t("📑  XLSX + CSV supported", "📑  支持 XLSX + CSV"))
    st.markdown(t("🧩  Map Excel letters", "🧩  按 Excel 栏位字母选择"))
    st.markdown(t("🧪  Duplicate-safe matching", "🧪  安全处理重复 Key"))
    st.divider()
    render_contact_cta(compact=True)
    st.divider()
    if st.button(t("↺  Reset workspace", "↺  清除当前工作空间"), use_container_width=True):
        reset_state()
        st.rerun()
    st.caption(t(
        "Per file: 10 MB · 30,000 rows · 100 columns. Files stay in this server session; choose your hosting carefully.",
        "每个文件：10 MB · 30,000 行 · 100 栏。文件保留在服务器会话内存中；请谨慎选择托管平台。",
    ))

hero_badge = t("NO DATABASE · USER-CONTROLLED", "无需数据库 · 由用户控制")
hero_main = t("Compare with <span>confidence.</span>", "让 Excel 对账<span>更简单。</span>")
hero_text = t(
    "Upload two Excel / CSV files, select the exact columns by letter, catch mismatches, and export a traceable report. No rigid header names. No formulas to write.",
    "上传两份 Excel / CSV，按 A、B、E、F 等栏位字母选择匹配与比较字段，发现差异、遗漏和重复数据，再下载可追溯到原始行号的报告。",
)
st.markdown(
    f'<div class="hero"><span class="pill">{hero_badge}</span><h1>{hero_main}</h1>'
    f'<p>{html.escape(hero_text)}</p><div class="hero-mini">'
    f'<span>{t("◉  Letter-based mapping", "◉  栏位字母选择")}</span>'
    f'<span>{t("◉  Duplicate detection", "◉  重复检测")}</span>'
    f'<span>{t("◉  Audit-ready XLSX", "◉  可追溯 Excel 报告")}</span>'
    '</div></div>', unsafe_allow_html=True,
)

section("01 / Source files", "01 / 上传文件", "Start with two spreadsheets", "上传两份 Excel / CSV",
        "Try sample data or use files you already have.", "可使用内建模拟数据，或上传你自己的文件。")
left_upload, right_upload = st.columns(2, gap="large")
with left_upload:
    st.markdown("**SOURCE A** &nbsp; <span class='cardtag'>Sales / Orders</span>", unsafe_allow_html=True)
    upload_a = st.file_uploader(t("Upload Excel / CSV A", "上传 Excel / CSV A"),
                                type=["xlsx", "csv"], key=f"up_a_{st.session_state.upload_epoch}")
with right_upload:
    st.markdown("**SOURCE B** &nbsp; <span class='cardtag'>Payments / Ledger</span>", unsafe_allow_html=True)
    upload_b = st.file_uploader(t("Upload Excel / CSV B", "上传 Excel / CSV B"),
                                type=["xlsx", "csv"], key=f"up_b_{st.session_state.upload_epoch}")
controls_a, controls_b, controls_note = st.columns([1.4, 1.8, 3])
with controls_a:
    if st.button(t("✦  Try sample", "✦  试用示例"), use_container_width=True):
        st.session_state.demo_mode = True
        st.session_state.report_signature = None
        st.rerun()
with controls_b:
    if st.session_state.demo_mode and upload_a is not None and upload_b is not None:
        if st.button(t("Use uploaded files", "切换为我的文件"), use_container_width=True):
            st.session_state.demo_mode = False
            st.session_state.report_signature = None
            st.rerun()
with controls_note:
    st.caption(t("Demo example: A matching B, amount E; B matching C, amount F.",
                 "示例：Excel A 用 B 栏匹配、E 栏比较；Excel B 用 C 栏匹配、F 栏比较。"))

sample_a, sample_b = demo_bytes()
with st.expander(t("↓ Download sample files", "↓ 下载示例 Excel 文件")):
    ca, cb = st.columns(2)
    with ca:
        st.download_button(t("Sample orders.xlsx", "示例订单.xlsx"), sample_a,
                           file_name="sample_orders.xlsx", use_container_width=True)
    with cb:
        st.download_button(t("Sample payments.xlsx", "示例付款.xlsx"), sample_b,
                           file_name="sample_payments.xlsx", use_container_width=True)

if st.session_state.demo_mode:
    data_a, name_a = sample_a, "sample_orders.xlsx"
    data_b, name_b = sample_b, "sample_payments.xlsx"
    st.success(t("Sample mode: all data is fictional. Upload two files to switch.",
                 "示例模式：所有数据均为虚构。上传两份文件后可切换为自己的数据。"))
elif upload_a is not None and upload_b is not None:
    data_a, name_a = upload_a.getvalue(), upload_a.name
    data_b, name_b = upload_b.getvalue(), upload_b.name
else:
    section("How it works", "工作流程", "From two files to clear answers", "三步完成对账",
            "Clear rules, transparent checks, useful results.", "规则由你决定，结果透明可追溯。")
    ca, cb, cc = st.columns(3)
    with ca:
        card("01 / Inspect", "01 / 检视", "Select two files and preview real rows and Excel column letters.",
             "选取两份文件，预览实际行及 Excel 栏位字母。")
    with cb:
        card("02 / Map", "02 / 选择", "Choose match keys, value columns and comparison rules.",
             "选择 Matching Key、比较字段及对账规则。")
    with cc:
        card("03 / Export", "03 / 输出", "Inspect exceptions and export a formatted, traceable workbook.",
             "检查异常并下载有原始行号的 Excel 报告。")
    st.info(t("Upload both files or click Try sample to begin.", "请上传两份文件，或点击「试用示例」开始。"))
    render_contact_cta()
    st.markdown('<div class="footer">ExcelDiff Pro · Python · Streamlit · OpenPyXL · No database</div>', unsafe_allow_html=True)
    st.stop()

try:
    sheets_a = list_sheets(data_a, name_a)
    sheets_b = list_sheets(data_b, name_b)
except DataImportError as exc:
    st.error(t("Cannot open the file: ", "无法开启文件：") + str(exc))
    st.stop()

section("02 / Inspect", "02 / 检视数据", "Choose worksheet and review sample rows", "选择 Worksheet 并预览数据",
        "A, B, C… refer to physical Excel positions, not assumed column names.",
        "A、B、C 等表示 Excel 实际栏位位置，而不是猜测的标题名称。")
sh_left, sh_right = st.columns(2, gap="large")
with sh_left:
    sheet_a = st.selectbox(t("Source A · Worksheet", "Excel A · 工作表"), sheets_a,
                           key=f"sheet_a_{short_hash(data_a)}")
with sh_right:
    sheet_b = st.selectbox(t("Source B · Worksheet", "Excel B · 工作表"), sheets_b,
                           key=f"sheet_b_{short_hash(data_b)}")
try:
    table_a = cache_table(data_a, name_a, sheet_a)
    table_b = cache_table(data_b, name_b, sheet_b)
except DataImportError as exc:
    st.error(t("Cannot read worksheet: ", "无法读取工作表：") + str(exc))
    st.stop()

with st.container(border=True):
    first, second = st.columns(2, gap="large")
    for holder, table, label in ((first, table_a, "A"), (second, table_b, "B")):
        with holder:
            st.markdown(f'**SOURCE {label}** &nbsp; <span class="cardtag">{html.escape(table.sheet)}</span>', unsafe_allow_html=True)
            st.caption(f"{table.height:,} {t('rows', '行')} · {table.width} {t('columns', '栏')} · {table.filename}")
            st.dataframe(raw_preview(table), use_container_width=True, hide_index=True, height=247)
            if table.width > 12:
                st.caption(t("Preview shows A–L; all columns remain selectable.",
                             "预览仅显示 A–L；其余栏位仍可在选单中选择。"))

# Import a mapping preset AFTER selecting sheets, BEFORE creating the mapping widgets.
with st.expander(t("Load a saved column mapping (.json)", "载入已保存的栏位 Mapping (.json)")):
    preset_file = st.file_uploader(t("Choose mapping profile", "选择 Mapping Profile"), type=["json"], key="mapping_file")
    if st.button(t("Apply saved mapping", "套用 Mapping 设置"), disabled=preset_file is None):
        try:
            loaded, saved_sheets = load_profile(preset_file.getvalue())
            st.session_state.preset_config = loaded
            st.session_state.preset_sheets = saved_sheets
            st.session_state.preset_epoch += 1
            st.session_state.report_signature = None
            st.rerun()
        except ValueError as exc:
            st.error(str(exc))

preset: CompareConfig = st.session_state.preset_config or CompareConfig()
if st.session_state.preset_sheets is not None:
    saved = st.session_state.preset_sheets
    if sheet_a != saved["a"] or sheet_b != saved["b"]:
        st.warning(t("The saved mapping was created for other worksheet names. Recheck EVERY column before comparing.",
                     "已保存的 Mapping 原本使用其他 Worksheet 名称，请重新确认每个栏位位置再执行比较。"))
    else:
        st.caption(t("Saved mapping profile loaded. Verify sample values before running.",
                     "已载入 Mapping Profile，请先核对实际 Sample Data。"))

# The preset epoch changes widget keys when a mapping profile is applied.
widget_prefix = f"{short_hash(data_a)}_{sheet_a}_{short_hash(data_b)}_{sheet_b}_{st.session_state.preset_epoch}"

def wk(identifier: str) -> str:
    return widget_prefix + ":" + identifier

section("03 / Rules", "03 / Mapping 设置", "Select your exact Excel columns", "选择 Excel 实际栏位字母",
        "Choose matching keys and values separately for A and B. Column headers may be different.",
        "Excel A 和 B 的 Matching Key、比较数值栏位可以不同，Header 名称也无需一致。")
start_left, start_right = st.columns(2, gap="large")
with start_left:
    start_a = int(st.number_input(t("A · First data row", "A · 数据开始行"), min_value=1, max_value=table_a.height,
                                  value=min(max(preset.start_a, 1), table_a.height), step=1, key=wk("start_a")))
with start_right:
    start_b = int(st.number_input(t("B · First data row", "B · 数据开始行"), min_value=1, max_value=table_b.height,
                                  value=min(max(preset.start_b, 1), table_b.height), step=1, key=wk("start_b")))
st.caption(t("Row 1 is usually a header, so select Row 2. If there is no header, select Row 1.",
             "如第 1 行是标题，一般选择 Row 2；没有标题则选择 Row 1。"))
key_left, key_right = st.columns(2, gap="large")
with key_left:
    key_a = choose_column(table_a, t("A · Matching Key", "A · 匹配编号栏位"), preset.keys_a[0], wk("key_a"), start_a)
with key_right:
    key_b = choose_column(table_b, t("B · Matching Key", "B · 匹配编号栏位"), preset.keys_b[0], wk("key_b"), start_b)

with st.expander(t("Advanced: match using two columns", "进阶：使用两个栏位一起匹配")):
    exl, exr = st.columns(2)
    with exl:
        labels_a = column_options(table_a, start_a)
        opts_a = [None] + [c for c in labels_a if c != key_a]
        default_a = preset.keys_a[1] if len(preset.keys_a) > 1 and preset.keys_a[1] in opts_a else None
        extra_key_a = st.selectbox(t("A · Second key", "A · 第二匹配栏位"), opts_a,
                                   index=opts_a.index(default_a),
                                   format_func=lambda c: t("None", "不使用") if c is None else labels_a[c], key=wk("extra_key_a"))
    with exr:
        labels_b = column_options(table_b, start_b)
        opts_b = [None] + [c for c in labels_b if c != key_b]
        default_b = preset.keys_b[1] if len(preset.keys_b) > 1 and preset.keys_b[1] in opts_b else None
        extra_key_b = st.selectbox(t("B · Second key", "B · 第二匹配栏位"), opts_b,
                                   index=opts_b.index(default_b),
                                   format_func=lambda c: t("None", "不使用") if c is None else labels_b[c], key=wk("extra_key_b"))
    st.caption(t("Example: Order ID + Currency. Choose both second keys, or neither.",
                 "例如 Order ID + Currency，请在 A 和 B 都选择第二个 Key，或都不选择。"))

modes = {"number": t("Compare amount / numbers", "比较金额 / 数字"),
         "text": t("Compare text / status", "比较文字 / 状态"),
         "presence": t("Check whether keys exist", "只检查编号是否存在")}
method = st.radio(t("Comparison method", "比较方式"), list(modes),
                  index=list(modes).index(preset.compare_type),
                  format_func=lambda k: modes[k], horizontal=True, key=wk("method"))
compare_a = compare_b = None
if method != "presence":
    val_left, val_right = st.columns(2, gap="large")
    with val_left:
        compare_a = choose_column(table_a, t("A · Value to compare", "A · 要比较的栏位"),
                                  preset.value_a if preset.value_a is not None else 4,
                                  wk("value_a"), start_a,
                                  exclude=(key_a,) + ((extra_key_a,) if extra_key_a is not None else ()))
    with val_right:
        compare_b = choose_column(table_b, t("B · Value to compare", "B · 要比较的栏位"),
                                  preset.value_b if preset.value_b is not None else 5,
                                  wk("value_b"), start_b,
                                  exclude=(key_b,) + ((extra_key_b,) if extra_key_b is not None else ()))

keys_a = (key_a,) + ((extra_key_a,) if extra_key_a is not None else ())
keys_b = (key_b,) + ((extra_key_b,) if extra_key_b is not None else ())

with st.expander(t("Advanced: output columns & matching rules", "进阶：输出栏位与比较规则"), expanded=False):
    additional_left, additional_right = st.columns(2, gap="large")
    with additional_left:
        labels_a = column_options(table_a, start_a)
        defaults_a = [c for c in preset.extra_a if c in labels_a]
        extra_a = tuple(st.multiselect(t("Other A columns to include in report", "报告中附加显示 A 的栏位"),
                                       list(labels_a), default=defaults_a,
                                       format_func=lambda c: labels_a[c], key=wk("extra_a")))
    with additional_right:
        labels_b = column_options(table_b, start_b)
        defaults_b = [c for c in preset.extra_b if c in labels_b]
        extra_b = tuple(st.multiselect(t("Other B columns to include in report", "报告中附加显示 B 的栏位"),
                                       list(labels_b), default=defaults_b,
                                       format_func=lambda c: labels_b[c], key=wk("extra_b")))
    trim_keys = st.checkbox(t("Trim spaces around matching IDs", "去掉匹配编号前后空格"),
                            value=preset.trim_keys, key=wk("trim_keys"))
    ignore_key_case = st.checkbox(t("Ignore uppercase / lowercase for matching IDs", "编号匹配时忽略大小写"),
                                  value=preset.ignore_key_case, key=wk("ignore_key_case"))
    trim_text = st.checkbox(t("Trim compared text", "比较文字前去掉前后空格"), value=preset.trim_text,
                            key=wk("trim_text"))
    ignore_text_case = st.checkbox(t("Ignore case in compared text", "比较文字时忽略大小写"),
                                   value=preset.ignore_text_case, key=wk("ignore_text_case"))
    if method == "number":
        fmt_left, fmt_right = st.columns(2)
        options = [".", ","]
        with fmt_left:
            decimal_a = st.selectbox(t("A · Number style", "A · 数字格式"), options,
                                     index=options.index(preset.decimal_a),
                                     format_func=lambda v: "1,234.56" if v == "." else "1.234,56",
                                     key=wk("decimal_a"))
        with fmt_right:
            decimal_b = st.selectbox(t("B · Number style", "B · 数字格式"), options,
                                     index=options.index(preset.decimal_b),
                                     format_func=lambda v: "1,234.56" if v == "." else "1.234,56",
                                     key=wk("decimal_b"))
        tolerance_text = st.text_input(t("Allowed amount difference", "允许的金额差额"),
                                       value=str(preset.tolerance), key=wk("tolerance"),
                                       help=t("Example: 0.01 means differences ≤ 0.01 count as matched.",
                                              "例如设置 0.01，代表差额不超过 0.01 就算一致。"))
        aggregate = st.checkbox(t("Group & SUM multiple payments for the same key", "相同编号的多笔付款合计比较"),
                                value=preset.aggregate, key=wk("aggregate"),
                                help=t("Enable only if repeated rows genuinely need to be summed.",
                                       "只有业务规则确实需要把多笔付款加总时才启用。"))
    else:
        decimal_a, decimal_b, tolerance_text, aggregate = ".", ".", "0", False

try:
    tolerance = Decimal(tolerance_text)
except (InvalidOperation, ValueError):
    tolerance = Decimal("-1")
configuration = CompareConfig(
    keys_a=keys_a, keys_b=keys_b, value_a=compare_a, value_b=compare_b,
    start_a=start_a, start_b=start_b, compare_type=method,
    decimal_a=decimal_a, decimal_b=decimal_b, tolerance=tolerance,
    trim_keys=trim_keys, ignore_key_case=ignore_key_case,
    trim_text=trim_text, ignore_text_case=ignore_text_case,
    aggregate=aggregate, extra_a=extra_a, extra_b=extra_b,
)
valid = True
if (extra_key_a is None) != (extra_key_b is None):
    st.error(t("Choose a second matching key in BOTH sources, or in neither.",
               "第二个 Key 需要在 A、B 两边都设置，或都不设置。"))
    valid = False
if tolerance < 0 or not tolerance.is_finite():
    st.error(t("Amount tolerance must be a non-negative finite number.",
               "允许差额必须是大于或等于 0 的有效数字。"))
    valid = False

if valid:
    profile_bytes = export_profile(configuration, sheet_a, sheet_b)
    st.download_button(t("↓ Save these settings as a mapping profile (.json)",
                         "↓ 保存这次的栏位 Mapping Profile (.json)"),
                       profile_bytes, file_name="exceldiff_mapping.json",
                       mime="application/json", use_container_width=True)
st.caption(t("Saved profiles store only column positions and rules — never your spreadsheet contents.",
             "Mapping Profile 只记录栏位位置和比较规则，不包含原始 Excel 数据。"))

section("04 / Check mapping", "04 / 检查 Mapping", "Inspect before comparison", "执行比较前先检查",
        "These checks catch likely mistakes, but you must confirm business meaning.",
        "系统可发现常见选择错误，但最终仍需要由用户确认所选栏位的业务含义。")
if valid:
    try:
        analysis = mapping_diagnostics(table_a, table_b, configuration)
        diagnosis_a = inspect_table(table_a, keys_a, start_a, configuration)
        diagnosis_b = inspect_table(table_b, keys_b, start_b, configuration)
    except ValueError as exc:
        st.error(str(exc))
        valid = False

if valid:
    diag_left, diag_right = st.columns(2, gap="large")
    for holder, tbl, keys, start, stats, label in (
        (diag_left, table_a, keys_a, start_a, diagnosis_a, "A"),
        (diag_right, table_b, keys_b, start_b, diagnosis_b, "B"),
    ):
        with holder:
            with st.container(border=True):
                st.markdown(f"**{t('Source', '文件')} {label} · {t('Key diagnostics','编号检查')}**")
                n1, n2, n3 = st.columns(3)
                n1.metric(t("Data rows", "数据行数"), stats["data_rows"])
                n2.metric(t("Empty keys", "空编号"), stats["empty_keys"])
                n3.metric(t("Duplicate keys", "重复编号"), stats["duplicate_keys"])
                st.caption(t("Selected key values from real rows:", "所选匹配栏位的实际数据样本："))
                st.dataframe(preview_keys(tbl, keys, start), use_container_width=True,
                             hide_index=True, height=172)
    percent_a = analysis["overlap_a_percent"]
    percent_b = analysis["overlap_b_percent"]
    left_count, right_count, overlap_count = st.columns(3)
    left_count.metric(t("Keys from A found in B", "A 编号在 B 中匹配率"), f"{percent_a:.1f}%")
    right_count.metric(t("Keys from B found in A", "B 编号在 A 中匹配率"), f"{percent_b:.1f}%")
    overlap_count.metric(t("Overlapping unique IDs", "两边共同编号数"), analysis["matched_unique_keys"])
    if analysis["matched_unique_keys"] == 0:
        st.error(t("No matching keys found. Check selected columns, row starts and key formatting before running.",
                   "没有任何相同编号！请检查栏位字母、起始 Row 和编号格式。"))
    elif percent_a < 25 or percent_b < 25:
        st.warning(t("Few keys match between the files. This may be expected — verify the mapping.",
                     "两个文件的匹配率偏低，可能是正常业务情况，也可能选错栏位，请确认。"))
    if method == "number" and (analysis["numeric_invalid_a"] or analysis["numeric_invalid_b"]):
        st.warning(t(
            f"Some amounts cannot be read as numbers (A: {analysis['numeric_invalid_a']}, B: {analysis['numeric_invalid_b']}). Check the selected columns and number styles.",
            f"部分值无法识别为金额（A: {analysis['numeric_invalid_a']}，B: {analysis['numeric_invalid_b']}）。请检查所选栏位和数字格式。",
        ))
    with st.expander(t("Preview matching records from both files", "预览两份文件中能配对的记录"), expanded=True):
        if analysis["preview"]:
            df = pd.DataFrame(analysis["preview"])
            if zh:
                df = df.rename(columns={"Key": "编号", "A row": "A 原始行", "A value": "A 数值", "B row": "B 原始行", "B value": "B 数值"})
            st.dataframe(df, use_container_width=True, hide_index=True)
        else:
            st.info(t("No matching record pairs to preview.", "目前没有可以预览的匹配记录。"))

signature = hashlib.sha256(repr((short_hash(data_a), short_hash(data_b), name_a, name_b,
                                 sheet_a, sheet_b, configuration)).encode()).hexdigest()
run_area, note_area = st.columns([1.5, 3.5])
with run_area:
    run = st.button(t("⟶  Run comparison", "⟶  开始比较"), type="primary",
                    use_container_width=True, disabled=not valid)
with note_area:
    st.caption(t("Results are updated only when you run again after changing the rules.",
                 "修改比较设置之后，需要重新点击「开始比较」才会更新报告。"))
if run:
    try:
        with st.spinner(t("Comparing records and generating report…", "正在比较数据并生成报告……")):
            outcome = compare_tables(table_a, table_b, configuration)
            report = export_report(table_a, table_b, outcome, configuration)
        st.session_state.report_signature = signature
        st.session_state.report_result = outcome
        st.session_state.report_blob = report
    except (ValueError, DataImportError) as exc:
        st.error(t("Unable to complete comparison: ", "无法完成比较：") + str(exc))
    except Exception:
        st.error(t("Report generation failed unexpectedly. Please verify the files.",
                   "产生报告时遇到异常，请检查文件和数据。"))

if st.session_state.report_result is not None and st.session_state.report_signature != signature:
    st.info(t("Your mapping changed. Run the comparison again to refresh the report.",
              "Mapping 设置已经改变，请重新执行比较以更新结果。"))

if st.session_state.report_result is not None and st.session_state.report_signature == signature:
    outcome = st.session_state.report_result
    metrics = outcome.metrics
    section("05 / Results", "05 / 比较结果", "Reconciliation is complete", "Excel 比较完成",
            "Find matched rows, discrepancies, missing records and duplicates — with original row numbers.",
            "查看一致、不一致、仅存在其中一边、重复和错误记录，所有记录均可追溯到原始 Excel Row。")
    k1, k2, k3, k4 = st.columns(4)
    k1.metric(t("Matched", "一致"), f"{metrics['matched']:,}")
    k2.metric(t("Different", "不一致"), f"{metrics['different']:,}")
    k3.metric(t("Only in A / B", "仅在 A / B"), f"{metrics['only_a'] + metrics['only_b']:,}")
    k4.metric(t("Needs review", "需要人工检查"), f"{metrics['duplicate'] + metrics['invalid_key'] + metrics['invalid_value']:,}")
    total = max(1, metrics["total_results"])
    breakdown = [(metrics["matched"], "#188B7D"), (metrics["different"], "#E5A66A"),
                 (metrics["only_a"] + metrics["only_b"], "#9381CE"),
                 (metrics["duplicate"] + metrics["invalid_key"] + metrics["invalid_value"], "#D66F6D")]
    width_bars = "".join(f'<span style="width:{100 * count / total:.4f}%;background:{color};display:block"></span>'
                         for count, color in breakdown if count)
    st.markdown(f'<div class="statusbar">{width_bars}</div>', unsafe_allow_html=True)
    for message in outcome.warnings:
        st.warning(message)
    tab_all, tab_ex, tab_export = st.tabs([
        t("All results", "全部结果"), t("Exceptions", "异常记录"), t("Download report", "下载报告"),
    ])
    with tab_all:
        f1, f2 = st.columns(2)
        available = sorted({row["Result"] for row in outcome.rows})
        with f1:
            selected = st.multiselect(t("Show statuses", "筛选结果类型"), available, default=available,
                                      key="result_status_filter")
        with f2:
            query = st.text_input(t("Search matching key", "查找匹配编号"), key="key_filter",
                                  placeholder="A002").strip().casefold()
        visible = [r for r in outcome.rows if r["Result"] in selected and query in str(r["Key"]).casefold()]
        st.caption(t(f"Showing {len(visible)} of {len(outcome.rows)} rows.",
                     f"显示 {len(visible)} / {len(outcome.rows)} 条结果。"))
        if visible:
            st.dataframe(result_table(visible), use_container_width=True, hide_index=True, height=370)
        else:
            st.info(t("No records match the filter.", "没有符合筛选条件的记录。"))
    with tab_ex:
        st.caption(t("Duplicates, missing keys and invalid amounts are not silently counted as matches.",
                     "重复编号、空 Key 和无效金额不会被悄悄判定为成功匹配。"))
        if outcome.exceptions:
            st.dataframe(result_table(outcome.exceptions), use_container_width=True,
                         hide_index=True, height=360)
        else:
            st.success(t("No exceptions found.", "没有发现异常记录。"))
    with tab_export:
        st.markdown("**" + t("Your comparison is ready to download", "报告可以下载了") + "**")
        st.caption(t("Excel report includes OVERVIEW, ALL RESULTS, EXCEPTIONS and SETTINGS sheets. Original files are not embedded.",
                     "Excel 报告包含 OVERVIEW、ALL RESULTS、EXCEPTIONS、SETTINGS，且不会附上原始完整文件。"))
        st.download_button(t("↓  Download Excel report", "↓  下载完整 Excel 比较报告"),
                           data=st.session_state.report_blob,
                           file_name="ExcelDiff_Comparison_Report.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           type="primary", use_container_width=True)
        st.download_button(t("↓  Download exceptions as CSV", "↓  下载异常记录 CSV"),
                           data=csv_export(outcome.exceptions),
                           file_name="ExcelDiff_Exceptions.csv", mime="text/csv",
                           use_container_width=True)

with st.expander(t("Help / FAQ and limitations", "说明 / 常见问题及限制")):
    if zh:
        st.markdown("""
**这是一个什么系统？** ExcelDiff Pro 是无需数据库的两份 Excel / CSV 比较和对账系统。

**什么是 Matching Key？** 用来认出同一笔记录的编号。例如：Excel A 的 B 栏订单号，对应 Excel B 的 C 栏 Reference。

**如何避免选错栏位？** 每个选单同时显示栏位字母、靠近起始行的标题与数据样本；执行前会检查空 Key、重复、匹配率与金额格式。系统不能自动保证你选择的是正确业务字段。

**重复订单怎么办？** 默认把相同 Key 的多行列为 DUPLICATE。只有确实需要多笔付款加总时，才启用 Group & SUM。

**我的数据安全吗？** 这个 App 不使用数据库、不调用 AI API，也不会主动将资料储存在磁盘。然而上传的文件需要传输到托管服务器并暂存在内存；服务器及运营商可能保留流量日志。机密数据建议内部私有部署。

**支持所有 Excel 格式吗？** 不支持。仅支持表格式 XLSX 与 CSV；旧版 XLS、带密码、复杂合并格、交叉表、宏与需要计算但未缓存的公式，都可能需要预处理。第一版每次比较一组数值栏位，也能选多个 Matching Key。

**Column B 下次变成 C 怎么办？** 应重新检查 Preview。Mapping Profile 记录的是栏位位置，原文件插入栏位后需要更新。CSV 的栏位也按 A、B、C 顺序编号。
""")
    else:
        st.markdown("""
**What is ExcelDiff Pro?** A database-free Excel / CSV comparison and reconciliation web app.

**What is a matching key?** An identifier such as Order ID (B in file A) matched to Reference (C in file B).

**How do I avoid incorrect columns?** Each choice displays the Excel letter, a nearby header and real sample values. Pre-flight checks inspect blanks, duplicates, overlap and numeric formats. These checks cannot infer your business intent.

**What about duplicate orders?** Duplicate keys are flagged by default. Enable Group & SUM only when summing repeated payments is the correct business rule.

**Are uploads secure?** The app uses no database, third-party AI or analytics. Files are still uploaded to the host server's memory; the host may maintain operational logs. Deploy privately for sensitive spreadsheets.

**Does it support every Excel?** No. The supported formats are tabular XLSX / CSV. Password-protected, merged/pivoted layouts, macros and formulas without saved results need preprocessing. The first release compares one value-column pair, with optional composite match keys.

**What if column B becomes C next month?** Always inspect the preview. Mapping profiles store physical column locations, not permanent semantic meanings.
""")

render_contact_cta()
st.markdown('<div class="footer">ExcelDiff Pro · Python + Streamlit + OpenPyXL · No database · No external AI calls</div>', unsafe_allow_html=True)

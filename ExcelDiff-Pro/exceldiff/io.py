"""Safe-ish, bounded, in-memory readers for ExcelDiff Pro (XLSX and CSV)."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 30_000
MAX_COLUMNS = 100
MAX_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
MAX_ZIP_MEMBERS = 500


class DataImportError(ValueError):
    """A user-friendly unsupported or unsafe file error."""


@dataclass(frozen=True)
class DataTable:
    filename: str
    sheet: str
    rows: tuple[tuple[Any, ...], ...]
    width: int

    @property
    def height(self) -> int:
        return len(self.rows)

    def get(self, excel_row: int, zero_based_col: int) -> Any:
        if excel_row < 1 or excel_row > self.height:
            return None
        row = self.rows[excel_row - 1]
        return row[zero_based_col] if zero_based_col < len(row) else None

    def data_rows(self, start_row: int) -> list[tuple[int, tuple[Any, ...]]]:
        if start_row < 1:
            raise ValueError("Starting row must be at least 1")
        return [
            (idx, row)
            for idx, row in enumerate(self.rows[start_row - 1 :], start=start_row)
            if any(not is_blank(cell) for cell in row)
        ]


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        import math
        if isinstance(value, float) and math.isnan(value):
            return True
    except Exception:
        pass
    return False


def present(value: Any, max_len: int = 85) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ", timespec="seconds")
    elif isinstance(value, (date, time)):
        text = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.replace("\r", " ").replace("\n", " ")
    return text if len(text) <= max_len else text[: max_len - 1] + "…"


def check_file(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in (".xlsx", ".csv"):
        raise DataImportError("Only .xlsx and .csv files are supported. Convert .xls or .xlsm to .xlsx first.")
    if not data:
        raise DataImportError("This file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise DataImportError("This file exceeds the 10 MB per-file limit.")
    if ext == ".xlsx":
        if not data.startswith(b"PK"):
            raise DataImportError("This does not appear to be a valid XLSX workbook.")
        try:
            with ZipFile(BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > MAX_ZIP_MEMBERS:
                    raise DataImportError("The workbook contains too many internal entries.")
                total = sum(info.file_size for info in infos)
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise DataImportError("The workbook expands beyond the 120 MB safety limit.")
                if any(info.flag_bits & 0x1 for info in infos):
                    raise DataImportError("Password-protected Excel workbooks are not supported.")
        except BadZipFile as exc:
            raise DataImportError("Invalid or damaged XLSX file.") from exc
    return ext


def list_sheets(data: bytes, filename: str) -> list[str]:
    ext = check_file(data, filename)
    if ext == ".csv":
        return ["CSV Data"]
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception as exc:
        raise DataImportError("Could not open this Excel workbook. Check that it is valid and not encrypted.") from exc


def _decode_csv(data: bytes) -> str:
    encodings = ["utf-8-sig", "gb18030", "cp1252"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ["utf-16"]
    for encoding in encodings:
        try:
            result = data.decode(encoding, errors="strict")
            if "\x00" in result:
                continue
            return result
        except (UnicodeError, ValueError):
            continue
    raise DataImportError("Could not decode CSV. Save it as UTF-8 CSV and retry.")


def _read_csv(data: bytes, filename: str) -> DataTable:
    text = _decode_csv(data)
    try:
        sample = text[:32768]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(StringIO(text, newline=""), dialect)
        records: list[tuple[Any, ...]] = []
        width = 0
        for row in reader:
            if len(records) >= MAX_ROWS:
                raise DataImportError(f"Only up to {MAX_ROWS:,} physical rows per file are supported.")
            if len(row) > MAX_COLUMNS:
                raise DataImportError(f"Only up to {MAX_COLUMNS} columns per file are supported.")
            width = max(width, len(row))
            records.append(tuple(row))
    except (csv.Error, UnicodeError) as exc:
        raise DataImportError("Could not parse CSV. Check the delimiter and quoting.") from exc
    if not records or width == 0:
        raise DataImportError("No data was found in this CSV file.")
    padded = tuple(row + (None,) * (width - len(row)) for row in records)
    return DataTable(filename=filename, sheet="CSV Data", rows=padded, width=width)


def _read_xlsx(data: bytes, filename: str, sheet: str | None) -> DataTable:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheet_name = sheet if sheet is not None else workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise DataImportError(f"Worksheet '{sheet_name}' was not found.")
            ws = workbook[sheet_name]
            if ws.max_column and ws.max_column > MAX_COLUMNS:
                raise DataImportError(f"This sheet exceeds the {MAX_COLUMNS}-column limit.")
            if ws.max_row and ws.max_row > MAX_ROWS:
                raise DataImportError(f"This sheet exceeds the {MAX_ROWS:,}-row limit.")
            width = max(1, ws.max_column or 1)
            records: list[tuple[Any, ...]] = []
            for row in ws.iter_rows(max_col=width, values_only=True):
                if len(records) >= MAX_ROWS:
                    raise DataImportError(f"Only up to {MAX_ROWS:,} rows are supported.")
                records.append(tuple(row))
            if not records or all(is_blank(v) for row in records for v in row):
                raise DataImportError("The selected sheet has no data.")
            return DataTable(filename=filename, sheet=sheet_name, rows=tuple(records), width=width)
        finally:
            workbook.close()
    except DataImportError:
        raise
    except Exception as exc:
        raise DataImportError("Could not read the worksheet. Try exporting it as a fresh .xlsx file.") from exc


def load_table(data: bytes, filename: str, sheet: str | None = None) -> DataTable:
    ext = check_file(data, filename)
    if ext == ".csv":
        return _read_csv(data, filename)
    return _read_xlsx(data, filename, sheet)


def column_options(table: DataTable, start_row: int = 2) -> dict[int, str]:
    results = {}
    for col in range(table.width):
        letter = get_column_letter(col + 1)
        header = present(table.get(start_row - 1, col), 26) if start_row > 1 else ""
        samples = []
        for row_num in range(start_row, min(table.height, start_row + 18) + 1):
            value = present(table.get(row_num, col), 24)
            if value and value not in samples:
                samples.append(value)
            if len(samples) >= 2:
                break
        label = f"{letter}"
        if header:
            label += f"  ·  {header}"
        if samples:
            label += "    |    " + " / ".join(samples)
        results[col] = label
    return results

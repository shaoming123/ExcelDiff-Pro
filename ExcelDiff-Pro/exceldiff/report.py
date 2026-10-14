"""Style-rich Excel reports; only writes in memory, never to the server disk."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

from .core import (
    ComparisonResult, CompareConfig, MATCHED, DIFFERENT, ONLY_A, ONLY_B,
    DUPLICATE, INVALID_KEY_A, INVALID_KEY_B, INVALID_VALUE,
)
from .io import DataTable, present

NAVY = "12233D"
INK = "1F2F44"
SUBTLE = "61718A"
TEAL = "078878"
WHITE = "FFFFFF"
PALE = "EFF4F8"
LINE = "DEE5EF"
STATUS_STYLES = {
    MATCHED: ("E7F6EE", "187847"),
    DIFFERENT: ("FFF1E8", "B35116"),
    ONLY_A: ("F0EEFF", "6553B7"),
    ONLY_B: ("F0EEFF", "6553B7"),
    DUPLICATE: ("FFF3D5", "906313"),
    INVALID_VALUE: ("FCE8E8", "AE3232"),
    INVALID_KEY_A: ("FCE8E8", "AE3232"),
    INVALID_KEY_B: ("FCE8E8", "AE3232"),
}


def safe_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        # An output workbook must never contain user-supplied executable formulas.
        value = value.replace("\x00", "")[:32000]
        if value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def write_safe(ws, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row, col, safe_cell(value))
    if isinstance(cell.value, str):
        cell.data_type = "s"


def title_bar(ws, heading: str, subtitle: str, width: int) -> None:
    width = max(width, 7)
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=width)
    top = ws.cell(1, 1, heading)
    top.font = Font(name="Aptos Display", size=21, bold=True, color=WHITE)
    top.fill = PatternFill("solid", fgColor=NAVY)
    top.alignment = Alignment(vertical="center", indent=1)
    for line in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=width):
        for cell in line:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
    sub = ws.cell(3, 1, subtitle)
    sub.font = Font(name="Aptos", size=10, italic=True, color=SUBTLE)
    sub.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[3].height = 28


def _add_results_sheet(workbook: Workbook, title: str, items: list[dict[str, Any]], *, all_headers: list[str]):
    ws = workbook.create_sheet(title)
    width = len(all_headers)
    title_bar(ws, title, "EXCELDIFF PRO · Deterministic reconciliation results", width)
    header_row = 5
    for idx, name in enumerate(all_headers, start=1):
        cell = ws.cell(header_row, idx, name)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name="Aptos", size=10, color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 28
    for i, record in enumerate(items, start=header_row + 1):
        for col, name in enumerate(all_headers, start=1):
            write_safe(ws, i, col, record.get(name, ""))
            cell = ws.cell(i, col)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F6F9FC")
            if name == "Result" and record.get("Result") in STATUS_STYLES:
                bg, fg = STATUS_STYLES[record["Result"]]
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(name="Aptos", size=10, bold=True, color=fg)
        ws.row_dimensions[i].height = 23
    ws.freeze_panes = "C6"
    if items:
        ws.auto_filter.ref = f"A5:{get_column_letter(width)}{header_row + len(items)}"
    for col_idx, header in enumerate(all_headers, start=1):
        if header == "Key":
            column_width = 24
        elif header == "Result":
            column_width = 22
        elif "rows" in header.lower():
            column_width = 18
        elif "Value" in header or "Difference" in header:
            column_width = 20
        else:
            column_width = 25
        ws.column_dimensions[get_column_letter(col_idx)].width = column_width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_options.horizontalCentered = True
    return ws


def export_report(a: DataTable, b: DataTable, result: ComparisonResult, cfg: CompareConfig) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "OVERVIEW"
    title_bar(overview, "ExcelDiff Pro", "COMPARISON REPORT  /  Data stays out of your database", 7)

    overview["A5"] = "RECONCILIATION SUMMARY"
    overview["A5"].font = Font(name="Aptos", size=12, bold=True, color=TEAL)
    metrics = [
        ("Source A rows", result.metrics["file_a_rows"]),
        ("Source B rows", result.metrics["file_b_rows"]),
        ("Unique keys evaluated", result.metrics["keys_checked"]),
        ("Matched", result.metrics["matched"]),
        ("Different", result.metrics["different"]),
        ("Only in A", result.metrics["only_a"]),
        ("Only in B", result.metrics["only_b"]),
        ("Duplicate keys needing review", result.metrics["duplicate"]),
        ("Empty keys", result.metrics["invalid_key"]),
        ("Invalid comparison values", result.metrics["invalid_value"]),
    ]
    for row_idx, (label, count) in enumerate(metrics, start=7):
        overview.cell(row_idx, 1, label)
        overview.cell(row_idx, 4, count)
        for col in range(1, 7):
            c = overview.cell(row_idx, col)
            if row_idx % 2:
                c.fill = PatternFill("solid", fgColor=PALE)
            c.font = Font(name="Aptos", size=11, color=INK, bold=col == 4)
        overview.row_dimensions[row_idx].height = 23
    overview["A19"] = "HOW TO READ THIS REPORT"
    overview["A19"].font = Font(name="Aptos", size=12, color=TEAL, bold=True)
    notes = [
        "MATCHED = key found on both sides and selected values agree (within tolerance).",
        "DIFFERENT = key found on both sides but values differ.",
        "ONLY IN A/B = key absent from the other source.",
        "DUPLICATE KEY = repeated matching key; review it, or choose Group & SUM explicitly.",
        "EMPTY KEY = matching key is blank, so that row cannot be matched.",
        "INVALID VALUE = numeric text could not be parsed safely.",
        "Always verify column positions and sample records before relying on a comparison.",
        "Original input files are NOT included in this report.",
    ]
    for i, note in enumerate(notes, start=21):
        overview.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)
        cell = overview.cell(i, 1, note)
        cell.font = Font(name="Aptos", size=10, color=SUBTLE)
        overview.row_dimensions[i].height = 23
    overview.column_dimensions["A"].width = 37
    for letter in "BCDEFG":
        overview.column_dimensions[letter].width = 13
    overview.sheet_view.showGridLines = False
    overview.freeze_panes = "A5"

    headers = list(result.rows[0].keys()) if result.rows else ["Key", "Result", "A · Excel rows", "B · Excel rows", "A · Value", "B · Value", "Difference (A - B)"]
    _add_results_sheet(workbook, "ALL RESULTS", result.rows, all_headers=headers)
    _add_results_sheet(workbook, "EXCEPTIONS", result.exceptions, all_headers=headers)

    settings = workbook.create_sheet("SETTINGS")
    title_bar(settings, "Comparison settings", "Review these settings to reproduce the results", 7)
    from openpyxl.utils import get_column_letter
    info = [
        ("Created at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("File A", a.filename), ("File A sheet", a.sheet),
        ("File B", b.filename), ("File B sheet", b.sheet),
        ("Data starts at (A / B)", f"{cfg.start_a} / {cfg.start_b}"),
        ("Matching columns (A)", ", ".join(get_column_letter(c + 1) for c in cfg.keys_a)),
        ("Matching columns (B)", ", ".join(get_column_letter(c + 1) for c in cfg.keys_b)),
        ("Comparison mode", cfg.compare_type),
        ("Comparison A column", get_column_letter(cfg.value_a + 1) if cfg.value_a is not None and cfg.compare_type != "presence" else "—"),
        ("Comparison B column", get_column_letter(cfg.value_b + 1) if cfg.value_b is not None and cfg.compare_type != "presence" else "—"),
        ("Tolerance", str(cfg.tolerance) if cfg.compare_type == "number" else "N/A"),
        ("Decimal symbol A / B", f"{cfg.decimal_a} / {cfg.decimal_b}"),
        ("Ignore key case", str(cfg.ignore_key_case)),
        ("Trim keys", str(cfg.trim_keys)),
        ("Group & SUM", str(cfg.aggregate)),
        ("Extra A columns", ", ".join(get_column_letter(c + 1) for c in cfg.extra_a)),
        ("Extra B columns", ", ".join(get_column_letter(c + 1) for c in cfg.extra_b)),
    ]
    for i, (key, value) in enumerate(info, start=5):
        settings.cell(i, 1, key)
        write_safe(settings, i, 3, value)
        settings.cell(i, 1).font = Font(name="Aptos", size=10, bold=True, color=INK)
        settings.cell(i, 3).font = Font(name="Aptos", size=10, color=SUBTLE)
        settings.row_dimensions[i].height = 22
        if i % 2 == 0:
            for col in range(1, 7):
                settings.cell(i, col).fill = PatternFill("solid", fgColor=PALE)
    settings.column_dimensions["A"].width = 33
    settings.column_dimensions["B"].width = 5
    settings.column_dimensions["C"].width = 55
    settings.sheet_view.showGridLines = False

    workbook.active = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

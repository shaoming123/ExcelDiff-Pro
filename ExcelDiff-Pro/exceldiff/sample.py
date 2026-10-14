"""Completely synthetic, license-free demo files for portfolio visitors."""
from __future__ import annotations

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _demo_workbook(headers: list[str], rows: list[list[object]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for record in rows:
        ws.append(record)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="152A48")
        cell.font = Font(name="Aptos", size=11, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    for i in range(2, len(rows) + 2):
        ws.row_dimensions[i].height = 23
        if i % 2 == 0:
            for c in ws[i]:
                c.fill = PatternFill("solid", fgColor="F2F7FA")
    for idx, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(30, max(17, len(h) + 5))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def demo_bytes() -> tuple[bytes, bytes]:
    orders = _demo_workbook(
        ["No", "Order No", "Customer", "Order Date", "Total (RM)", "Channel"],
        [
            [1, "A001", "Ali", "2026-10-01", 100, "Online"],
            [2, "A002", "Tan", "2026-10-01", 250, "Retail"],
            [3, "A003", "Lee", "2026-10-02", 180, "Online"],
            [4, "A005", "Maya", "2026-10-03", 150, "Online"],
            [5, "A005", "Maya", "2026-10-03", 150, "Online"],
            [6, "A006", "Sara", "2026-10-03", "RM 75.00", "Retail"],
            [7, "A007", "Amir", "2026-10-04", 100, "Online"],
            [8, None, "No ID", "2026-10-04", 50, "Retail"],
        ],
        "Orders",
    )
    payments = _demo_workbook(
        ["Txn", "Method", "Reference", "Payment Date", "Status", "Paid Amount"],
        [
            [1001, "FPX", "A001", "2026-10-01", "Paid", 100],
            [1002, "Card", "A002", "2026-10-01", "Paid", 200],
            [1003, "Cash", "A004", "2026-10-02", "Paid", 90],
            [1004, "FPX", "A005", "2026-10-03", "Paid", 150],
            [1005, "Card", "A006", "2026-10-03", "Paid", "RM75.00"],
            [1006, "Card", "A007", "2026-10-04", "Paid", "pending"],
            [1007, "FPX", None, "2026-10-04", "Paid", 10],
        ],
        "Payments",
    )
    return orders, payments

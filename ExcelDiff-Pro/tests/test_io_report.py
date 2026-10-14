from io import BytesIO
import zipfile
from openpyxl import load_workbook
import pytest
from exceldiff.io import load_table, list_sheets, DataImportError, column_options
from exceldiff.core import CompareConfig, compare_tables
from exceldiff.report import export_report
from exceldiff.sample import demo_bytes


def test_sample_data_end_to_end():
    orders, payments=demo_bytes()
    assert list_sheets(orders,"orders.xlsx")==["Orders"]
    assert list_sheets(payments,"payments.xlsx")==["Payments"]
    a=load_table(orders,"orders.xlsx","Orders")
    b=load_table(payments,"payments.xlsx","Payments")
    assert a.width==6 and b.width==6
    assert column_options(a,2)[1].startswith("B")
    cfg=CompareConfig(keys_a=(1,),keys_b=(2,),value_a=4,value_b=5,start_a=2,start_b=2,extra_a=(2,),extra_b=(4,))
    res=compare_tables(a,b,cfg)
    assert res.metrics=={"file_a_rows":8,"file_b_rows":7,"keys_checked":7,"matched":2,"different":1,"only_a":1,"only_b":1,"duplicate":1,"invalid_key":2,"invalid_value":1,"total_results":9}
    report=export_report(a,b,res,cfg)
    wb=load_workbook(BytesIO(report),data_only=False)
    assert wb.sheetnames==["OVERVIEW","ALL RESULTS","EXCEPTIONS","SETTINGS"]
    assert wb["ALL RESULTS"].max_row==5+res.metrics["total_results"]
    assert wb["EXCEPTIONS"].max_row==5+len(res.exceptions)
    assert wb["OVERVIEW"]["D10"].value==2


def test_injection_is_escaped_in_excel_report():
    from exceldiff.io import DataTable
    a=DataTable("A.xlsx","X",(("k","v","note"),("ID1",1,"=HYPERLINK(\"https://evil.example\")")),3)
    b=DataTable("B.xlsx","X",(("k","v"),("ID1",1)),2)
    cfg=CompareConfig(keys_a=(0,),keys_b=(0,),value_a=1,value_b=1,start_a=2,start_b=2,extra_a=(2,))
    result=compare_tables(a,b,cfg)
    wb=load_workbook(BytesIO(export_report(a,b,result,cfg)))
    ws=wb["ALL RESULTS"]
    values=[x.value for x in ws[6]]
    suspicious=[x for x in values if isinstance(x,str) and 'HYPERLINK' in x]
    assert suspicious and suspicious[0].startswith("'=")
    assert all(x.data_type != 'f' for x in ws[6])


def test_csv_without_header_and_semicolon_delimiter():
    data=b"001;100\r\n002;200\r\n"
    t=load_table(data,"records.csv")
    assert t.width==2 and t.height==2 and t.get(1,0)=="001"


def test_csv_utf8_and_chinese_text():
    data="订单,总额\n001,200\n002,300\n".encode("utf-8-sig")
    t=load_table(data,"chinese.csv")
    assert t.get(1,0)=="订单"


def test_invalid_files_and_extensions():
    with pytest.raises(DataImportError):load_table(b"abc","test.exe")
    with pytest.raises(DataImportError):load_table(b"not-a-zip","test.xlsx")
    with pytest.raises(DataImportError):load_table(b"","test.csv")


def test_oversized_file_rejected():
    with pytest.raises(DataImportError):load_table(b"a"*(10*1024*1024+1),"test.csv")


def test_xlsx_zip_bomb_metadata_rejected():
    payload=b"0"*(121*1024*1024)
    out=BytesIO()
    with zipfile.ZipFile(out,mode="w",compression=zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        z.writestr("huge.xml",payload)
    with pytest.raises(DataImportError,match="expands"):
        load_table(out.getvalue(),"test.xlsx")

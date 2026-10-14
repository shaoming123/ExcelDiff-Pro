"""Lightweight no-browser route smoke tests for the Streamlit application.

These exercise Python control flow through the demonstration and result-export
paths without pretending to replace a real Streamlit/browser smoke test.
"""
from __future__ import annotations

import runpy
import sys
from types import ModuleType
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
import pytest


class StreamlitStopped(Exception):
    pass


class State(dict):
    def __getattr__(self, key):
        return self[key]
    def __setattr__(self, key, value):
        self[key] = value


class FakeUI(ModuleType):
    def __init__(self, language=None, demo_mode=True):
        super().__init__("streamlit")
        self.session_state=State(demo_mode=demo_mode)
        self.sidebar=self
        self.language=language
        self.rendered_html=[]
        self.injected_styles=[]

    def __enter__(self):
        return self
    def __exit__(self, exc_type, exc, tb):
        return False
    def __getattr__(self, name):
        if name in ("expander", "container", "spinner"):
            return lambda *a,**k:self
        if name == "html":
            return lambda body, **kwargs: self.injected_styles.append(body)
        if name == "markdown":
            return lambda body, **kwargs: self.rendered_html.append(str(body))
        if name in ("caption", "error", "warning", "success", "info", "metric",
                    "set_page_config", "download_button", "divider", "dataframe"):
            return lambda *a, **k: None
        raise AttributeError(name)
    def columns(self, count, **kw):
        return [self for _ in range(count if isinstance(count,int) else len(count))]
    def tabs(self, labels):
        return [self for _ in labels]
    def selectbox(self, label, options=None, index=0, **kw):
        if label=="Language / 语言":
            assert index == 0, "English must be the default language"
            return self.language if self.language is not None else options[index]
        return options[index]
    def radio(self, label, options, index=0, **kw):
        return options[index]
    def number_input(self, label, **kw):
        return kw['value']
    def multiselect(self, label, options=None, default=None, **kw):
        return default or []
    def text_input(self, label, **kw):
        return kw.get('value', '')
    def checkbox(self, label, value=False, **kw):
        return value
    def file_uploader(self, *a, **kw):
        return None
    def button(self, label, **kw):
        return ('Run comparison' in label or '开始比较' in label) and not kw.get('disabled',False)
    def stop(self):
        raise StreamlitStopped()
    def rerun(self):
        raise AssertionError('Unexpected st.rerun() during sample smoke test')


@pytest.mark.parametrize('language',['English','简体中文'])
def test_demo_app_completes_full_reconciliation(monkeypatch, language):
    fake=FakeUI(language)
    monkeypatch.setitem(sys.modules,"streamlit",fake)
    runpy.run_path(str(Path(__file__).parents[1]/'app.py'),run_name='__main__')
    result=fake.session_state.report_result
    assert len(fake.injected_styles) == 1
    assert Path(fake.injected_styles[0]).name == 'theme.css'
    assert any('greatfoodprefer@gmail.com' in text for text in fake.rendered_html)
    assert any('mailto:greatfoodprefer@gmail.com' in text for text in fake.rendered_html)
    assert result.metrics['matched']==2
    assert result.metrics['different']==1
    assert result.metrics['duplicate']==1
    assert result.metrics['invalid_value']==1
    exported=load_workbook(BytesIO(fake.session_state.report_blob), read_only=True)
    assert 'ALL RESULTS' in exported.sheetnames
    assert 'EXCEPTIONS' in exported.sheetnames


def test_homepage_english_default_contains_custom_service_cta(monkeypatch):
    fake=FakeUI(demo_mode=False)
    monkeypatch.setitem(sys.modules, "streamlit", fake)
    with pytest.raises(StreamlitStopped):
        runpy.run_path(str(Path(__file__).parents[1]/'app.py'),run_name='__main__')
    combined='\n'.join(fake.rendered_html)
    assert 'Need something beyond a simple spreadsheet comparison?' in combined
    assert 'multi-file reconciliation' in combined
    assert 'greatfoodprefer@gmail.com' in combined
    assert 'mailto:greatfoodprefer@gmail.com?subject=' in combined
    assert 'Let’s discuss your project' in combined
    assert '需要比基本 Excel 对比' not in combined


def test_homepage_chinese_contact_cta(monkeypatch):
    fake=FakeUI(language='简体中文', demo_mode=False)
    monkeypatch.setitem(sys.modules,"streamlit",fake)
    with pytest.raises(StreamlitStopped):
        runpy.run_path(str(Path(__file__).parents[1]/'app.py'),run_name='__main__')
    combined='\n'.join(fake.rendered_html)
    assert '需要比基本 Excel 对比更进一步的解决方案' in combined
    assert '邮件咨询定制项目' in combined
    assert 'greatfoodprefer@gmail.com' in combined


def test_stylesheet_contains_only_css_not_html():
    css=(Path(__file__).parents[1]/'assets'/'theme.css').read_text(encoding='utf-8')
    assert '<style>' not in css
    assert '</style>' not in css
    assert '[data-testid="stDataFrame"]' in css
    assert '.contact-panel {' in css

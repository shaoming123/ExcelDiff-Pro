from io import BytesIO
from openpyxl import load_workbook
from exceldiff.cli import run
from exceldiff.sample import demo_bytes
from exceldiff.core import CompareConfig
from exceldiff.presets import export_profile


def test_cli_generates_report_with_letter_mapping(tmp_path):
    a, b = demo_bytes()
    path_a, path_b = tmp_path / "sales.xlsx", tmp_path / "payments.xlsx"
    path_a.write_bytes(a)
    path_b.write_bytes(b)
    target = tmp_path / "out" / "reconciliation.xlsx"
    code=run([str(path_a),str(path_b),"--sheet-a","Orders","--sheet-b","Payments",
              "--key-a","B","--key-b","C","--value-a","E","--value-b","F",
              "--output",str(target)])
    assert code==0
    wb=load_workbook(target)
    assert wb.sheetnames==["OVERVIEW","ALL RESULTS","EXCEPTIONS","SETTINGS"]
    assert wb["OVERVIEW"]["D10"].value==2


def test_cli_works_with_mapping_preset(tmp_path):
    a, b = demo_bytes()
    (tmp_path/'a.xlsx').write_bytes(a)
    (tmp_path/'b.xlsx').write_bytes(b)
    profile=export_profile(CompareConfig(),"Orders","Payments")
    (tmp_path/'profile.json').write_bytes(profile)
    assert run([str(tmp_path/'a.xlsx'),str(tmp_path/'b.xlsx'),"--preset",str(tmp_path/'profile.json'),
                "--output",str(tmp_path/'result.xlsx')]) == 0
    assert (tmp_path/'result.xlsx').exists()


def test_cli_invalid_config_returns_error(tmp_path):
    assert run([str(tmp_path/'missing.xlsx'),str(tmp_path/'bad.xlsx'),"--key-a","A","--key-b","A",
                "--value-a","B","--value-b","B"]) == 2
